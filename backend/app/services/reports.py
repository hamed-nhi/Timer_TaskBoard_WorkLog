"""Read-only reports: timesheet, dashboard, tag summary, calendar, entry report, Excel export."""
from datetime import date, timedelta
import io

import jdatetime

from . import calendar_events
from ..core.errors import ApiError
from ..core.backup import fingerprint
from ..core.time import JALALI_MONTHS, PERSIAN_WEEKDAYS, g2j, j2g, jalali_month_length, now_local, today_jalali, today_local, week_start_of
from ..db.session import Db
from .tags import budget_info
from .timecalc import TASK_SELECT, effective_minutes, row_to_dict
from .user_settings import get_goals
from .validation import clean_jalali


def aligned_minutes(db, uid, year, month, upto_month, upto_day):
    """Minutes of finished work in a period, counted only up to a cut-off inside it:
    month given -> days 1..upto_day of that month; no month -> from the first day of the
    year up to (upto_month, upto_day). This is what makes "10 days of this month vs the
    previous month" compare the same number of days."""
    query = ("SELECT COALESCE(SUM(duration_minutes), 0) t FROM tasks "
             "WHERE user_id = ? AND duration_minutes IS NOT NULL AND jalali_year = ?")
    params = [uid, year]
    if month:
        query += " AND jalali_month = ? AND jalali_day <= ?"
        params += [month, upto_day]
    else:
        query += " AND (jalali_month < ? OR (jalali_month = ? AND jalali_day <= ?))"
        params += [upto_month, upto_month, upto_day]
    return db.execute(query, params).fetchone()["t"]


def build_comparison(db, uid, year, month, today):
    """Compare the selected period with the one before it over the *same stretch of days*.
    A running month is cut at today (10 days in -> days 1-10 of the previous month); a finished
    month or year is compared in full. Returns None for a future period or when no year is chosen."""
    if not year:
        return None
    if month:
        if (year, month) > (today.year, today.month):
            return None
        prev_year, prev_month = (year - 1, 12) if month == 1 else (year, month - 1)
        elapsed = today.day if (year, month) == (today.year, today.month) else jalali_month_length(year, month)
        days = min(elapsed, jalali_month_length(prev_year, prev_month))
        cur = aligned_minutes(db, uid, year, month, month, days)
        prev = aligned_minutes(db, uid, prev_year, prev_month, prev_month, days)
        mname, pname = JALALI_MONTHS[month - 1], JALALI_MONTHS[prev_month - 1]
        span = (lambda d, n: f"{d} {n}") if days == 1 else (lambda d, n: f"1 تا {d} {n}")
        cur_label, prev_label = span(days, mname), span(days, pname)
        label = "ماه قبل"
        kind = "month"
    else:
        if year > today.year:
            return None
        if year == today.year:
            cut_month, cut_day = today.month, today.day
            cur_label = f"1 فروردین تا {cut_day} {JALALI_MONTHS[cut_month - 1]} {year}"
            prev_label = f"1 فروردین تا {cut_day} {JALALI_MONTHS[cut_month - 1]} {year - 1}"
            label = "سال قبل (تا همین روز سال)"
        else:
            cut_month, cut_day = 12, 31
            cur_label, prev_label = f"کل سال {year}", f"کل سال {year - 1}"
            label = "سال قبل"
        days = None
        cur = aligned_minutes(db, uid, year, None, cut_month, cut_day)
        prev = aligned_minutes(db, uid, year - 1, None, cut_month, cut_day)
        kind = "year"
    diff = round(cur - prev, 1)
    return {
        "kind": kind, "label": label, "days": days, "current_label": cur_label, "prev_label": prev_label,
        "current_minutes": round(cur, 1), "prev_minutes": round(prev, 1), "diff_minutes": diff,
        "diff_percent": round(diff / prev * 100, 1) if prev else None,
    }


def timesheet(db: Db, user: dict, date_: str | None):
    """Weekly timesheet: one row per tag, one column per day (Saturday..Friday), in minutes.
    `?date=YYYY-MM-DD` is any day of the wanted week (default: today). Read-only."""
    try:
        anchor = date.fromisoformat(date_) if date_ else today_local()
    except ValueError:
        raise ApiError("تاریخ نامعتبر است")
    start = week_start_of(anchor)
    days = [start + timedelta(days=i) for i in range(7)]
    now = now_local()
    rows_db = db.execute(TASK_SELECT + " WHERE tasks.user_id = ? AND tasks.gregorian_date BETWEEN ? AND ?",
                         (user["id"], days[0], days[-1])).fetchall()
    index = {d: i for i, d in enumerate(days)}
    by_tag = {}
    for r in rows_db:
        key = r["tag_id"]
        entry = by_tag.setdefault(key, {"tag_id": key, "tag_name": r["tag_name"] or "بدون تگ", "minutes": [0.0] * 7})
        entry["minutes"][index[r["gregorian_date"]]] += effective_minutes(r, now)
    rows = []
    for entry in by_tag.values():
        entry["minutes"] = [round(m, 2) for m in entry["minutes"]]
        entry["total"] = round(sum(entry["minutes"]), 2)
        rows.append(entry)
    rows.sort(key=lambda e: (e["tag_id"] is None, -e["total"]))
    day_totals = [round(sum(e["minutes"][i] for e in rows), 2) for i in range(7)]
    out_days = []
    for i, d in enumerate(days):
        j = g2j(d)
        out_days.append({"gregorian": d.isoformat(), "weekday": PERSIAN_WEEKDAYS[i], "jalali_year": j.year,
                         "jalali_month": j.month, "jalali_day": j.day, "month_name": JALALI_MONTHS[j.month - 1],
                         "is_today": d == today_local()})
    return {"week_start": days[0].isoformat(), "week_end": days[-1].isoformat(), "days": out_days,
            "rows": rows, "day_totals": day_totals, "total": round(sum(day_totals), 2),
            "prev": (days[0] - timedelta(days=7)).isoformat(), "next": (days[0] + timedelta(days=7)).isoformat()}


def dashboard(db: Db, user: dict, year: int | None, month: int | None):
    uid = user["id"]
    query = TASK_SELECT + " WHERE tasks.user_id = ? AND tasks.duration_minutes IS NOT NULL"
    params = [uid]
    if year:
        query += " AND tasks.jalali_year = ?"
        params.append(year)
    if month:
        query += " AND tasks.jalali_month = ?"
        params.append(month)
    rows = db.execute(query, params).fetchall()

    total_minutes = sum(r["duration_minutes"] for r in rows)
    running_count = db.execute("SELECT COUNT(*) c FROM tasks WHERE user_id = ? AND status='running'", (uid,)).fetchone()["c"]
    paused_count = db.execute("SELECT COUNT(*) c FROM tasks WHERE user_id = ? AND status='paused'", (uid,)).fetchone()["c"]

    by_day = {}
    for r in rows:
        key = f"{r['jalali_year']}/{r['jalali_month']:02d}/{r['jalali_day']:02d}"
        by_day[key] = by_day.get(key, 0) + r["duration_minutes"]
    by_day_list = [{"date": k, "minutes": round(v, 1)} for k, v in sorted(by_day.items())]

    by_month = {}
    all_rows = db.execute(
        "SELECT jalali_year, jalali_month, duration_minutes FROM tasks WHERE user_id = ? AND duration_minutes IS NOT NULL"
        + (" AND jalali_year = ?" if year else ""), [uid, year] if year else [uid]).fetchall()
    for r in all_rows:
        key = f"{r['jalali_year']}/{r['jalali_month']:02d}"
        by_month[key] = by_month.get(key, 0) + r["duration_minutes"]
    by_month_list = [
        {"date": k, "month_name": JALALI_MONTHS[int(k.split('/')[1]) - 1], "minutes": round(v, 1)}
        for k, v in sorted(by_month.items())
    ]

    by_tag = {}
    for r in rows:
        key = r["tag_id"]
        entry = by_tag.setdefault(key, {"tag_id": key, "tag_name": r["tag_name"] or "بدون تگ", "minutes": 0, "count": 0})
        entry["minutes"] += r["duration_minutes"]
        entry["count"] += 1
    by_tag_list = sorted(({**e, "minutes": round(e["minutes"], 1)} for e in by_tag.values()), key=lambda x: -x["minutes"])

    return {
        "total_minutes": round(total_minutes, 1),
        "total_hours": round(total_minutes / 60, 2),
        "task_count": len(rows),
        "running_count": running_count,
        "paused_count": paused_count,
        "active_days": len(by_day_list),
        "days_in_month": jalali_month_length(year, month) if year and month else None,
        "avg_daily_minutes": round(total_minutes / len(by_day_list), 1) if by_day_list else 0,
        "best_day": max(by_day_list, key=lambda d: d["minutes"]) if by_day_list else None,
        "daily_goal_minutes": get_goals(db, uid)["daily_goal_minutes"],
        "by_day": by_day_list,
        "by_month": by_month_list,
        "by_tag": by_tag_list,
        "comparison": build_comparison(db, uid, year, month, today_jalali()),
    }


def tag_summary(db: Db, user: dict, tag: str, year: int | None, month: int | None):
    """Cumulative time for one tag (or for tasks without a tag, tag=none), over all time unless a
    year / month is given: totals, active days, per-month breakdown and the titles that took the most time."""
    uid = user["id"]
    if tag == "none":
        tag_info = {"id": None, "name": "بدون تگ"}
        budget = None
        where, params = "tasks.tag_id IS NULL", []
    else:
        try:
            tag_id = int(tag)
        except ValueError:
            raise ApiError("تگ نامعتبر است")
        row = db.execute("SELECT * FROM tags WHERE id = ? AND user_id = ?", (tag_id, uid)).fetchone()
        if not row:
            raise ApiError("تگ نامعتبر است", 404)
        tag_info = {"id": row["id"], "name": row["name"]}
        budget = budget_info(db, uid, row)
        where, params = "tasks.tag_id = ?", [tag_id]
    if year:
        where += " AND tasks.jalali_year = ?"
        params.append(year)
    if month:
        where += " AND tasks.jalali_month = ?"
        params.append(month)

    rows = db.execute(f"SELECT tasks.* FROM tasks WHERE tasks.user_id = ? AND {where} ORDER BY gregorian_date",
                      [uid] + params).fetchall()
    now = now_local()
    done = [r for r in rows if r["duration_minutes"] is not None]
    open_rows = [r for r in rows if r["status"] != "done"]

    total = sum(r["duration_minutes"] for r in done)
    days = {r["gregorian_date"] for r in done}
    by_month, by_title = {}, {}
    for r in done:
        key = f"{r['jalali_year']}/{r['jalali_month']:02d}"
        by_month[key] = by_month.get(key, 0) + r["duration_minutes"]
        entry = by_title.setdefault(r["title"], {"title": r["title"], "minutes": 0, "count": 0})
        entry["minutes"] += r["duration_minutes"]
        entry["count"] += 1

    def label(r):
        return f"{r['jalali_year']}/{r['jalali_month']:02d}/{r['jalali_day']:02d}"

    recent = sorted(done, key=lambda r: (r["gregorian_date"], r["start_time"] or ""), reverse=True)[:10]
    return {
        "tag": tag_info,
        "budget": budget if budget and budget["budget_minutes"] else None,
        "recent": [{"title": r["title"], "date": label(r), "minutes": round(r["duration_minutes"], 1)} for r in recent],
        "total_minutes": round(total, 1),
        "task_count": len(done),
        "day_count": len(days),
        "avg_per_day_minutes": round(total / len(days), 1) if days else 0,
        "first_date": label(done[0]) if done else None,
        "last_date": label(done[-1]) if done else None,
        "open_count": len(open_rows),
        "open_minutes": round(sum(effective_minutes(r, now) for r in open_rows), 1),
        "by_month": [
            {"date": k, "month_name": JALALI_MONTHS[int(k.split("/")[1]) - 1], "minutes": round(v, 1)}
            for k, v in sorted(by_month.items())
        ],
        "top_titles": sorted(({**e, "minutes": round(e["minutes"], 1)} for e in by_title.values()),
                             key=lambda e: -e["minutes"])[:8],
    }


def calendar_month(db: Db, user: dict, year: int | None, month: int | None):
    if year is None or month is None or not 1 <= month <= 12 or not 1 <= year <= 3000:
        raise ApiError("سال یا ماه نامعتبر است")
    year, month, _ = clean_jalali(year, month, 1)
    days_in_month = 29 if month == 12 and not jdatetime.date(year, 1, 1).isleap() else (31 if month <= 6 else 30)

    uid = user["id"]
    now = now_local()
    work = {}
    for r in db.execute("SELECT * FROM tasks WHERE user_id = ? AND jalali_year = ? AND jalali_month = ?",
                        (uid, year, month)):
        entry = work.setdefault(r["jalali_day"], {"minutes": 0.0, "tasks": 0})
        entry["minutes"] += effective_minutes(r, now)
        entry["tasks"] += 1
    planned = {
        r["jalali_day"]: r["c"] for r in db.execute(
            "SELECT jalali_day, COUNT(*) c FROM planned_tasks WHERE user_id = ? AND jalali_year = ? AND jalali_month = ? "
            "GROUP BY jalali_day", (uid, year, month))
    }

    first = j2g(year, month, 1)
    days = []
    for day in range(1, days_in_month + 1):
        g = first + timedelta(days=day - 1)
        hy, hm, hd, exact = calendar_events.hijri_date(g)
        events = calendar_events.events_for_day(g)
        w = work.get(day)
        days.append({
            "day": day,
            "gregorian": g.isoformat(),
            "weekday": (g.weekday() - 5) % 7,  # Saturday = 0 ... Friday = 6
            "hijri": {"year": hy, "month": hm, "day": hd, "month_name": calendar_events.HIJRI_MONTHS[hm - 1], "exact": exact},
            "holiday": any(e["holiday"] for e in events),
            "events": events,
            "minutes": round(w["minutes"], 1) if w else 0,
            "task_count": w["tasks"] if w else 0,
            "planned_count": planned.get(day, 0),
        })
    return {"jalali_year": year, "jalali_month": month, "month_name": JALALI_MONTHS[month - 1],
            "days_in_month": days_in_month, "first_weekday": days[0]["weekday"], "days": days}


def entry_report(db: Db, user: dict, year: int | None, month: int | None):
    query = "SELECT * FROM tasks WHERE user_id = ? AND duration_minutes IS NOT NULL"
    params = [user["id"]]
    if year:
        query += " AND jalali_year = ?"
        params.append(year)
    if month:
        query += " AND jalali_month = ?"
        params.append(month)
    days = {}
    for r in db.execute(query, params).fetchall():
        key = f"{r['jalali_year']}/{r['jalali_month']:02d}/{r['jalali_day']:02d}"
        d = days.setdefault(key, {"date": key, "total_tasks": 0, "entered_tasks": 0, "total_minutes": 0})
        d["total_tasks"] += 1
        d["total_minutes"] += r["duration_minutes"] or 0
        if r["entered_main_system"]:
            d["entered_tasks"] += 1
    result = []
    for key in sorted(days):
        d = days[key]
        d["total_minutes"] = round(d["total_minutes"], 1)
        d["all_entered"] = d["entered_tasks"] == d["total_tasks"]
        result.append(d)
    return result


def export_excel(db: Db, user: dict, year: int | None):
    uid = user["id"]
    if year:
        rows = db.execute(TASK_SELECT + " WHERE tasks.user_id = ? AND tasks.jalali_year = ? "
                          "ORDER BY tasks.jalali_month, tasks.jalali_day, tasks.start_time, tasks.id", (uid, year)).fetchall()
    else:
        rows = db.execute(TASK_SELECT + " WHERE tasks.user_id = ? ORDER BY tasks.jalali_year, tasks.jalali_month, "
                          "tasks.jalali_day, tasks.start_time, tasks.id", (uid,)).fetchall()
    return build_workbook(rows), f"timelog_u{uid}_{year or 'all'}_{fingerprint(rows)}.xlsx"


def build_workbook(rows) -> io.BytesIO:
    # imported here: only needed for exporting
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    months_present = sorted(set((r["jalali_year"], r["jalali_month"]) for r in rows))
    header_font = Font(name="B Nazanin", bold=True, size=12)
    cell_font = Font(name="B Nazanin", size=11)
    header_fill = PatternFill(start_color="D9E1F2", end_color="D9E1F2", fill_type="solid")

    for (y, m) in months_present:
        ws = wb.create_sheet(f"{JALALI_MONTHS[m - 1]} {y}"[:31])
        ws.sheet_view.rightToLeft = True
        ws.append(["تاریخ (ماه/روز)", "شرح فعالیت", "تگ", "ساعت شروع", "ساعت پایان", "مدت (دقیقه)",
                   "نوع گزارش", "نام پروژه", "نام استوری", "تکمیلی (دقیقه)", "پلن شده", "خارج از شرکت"])
        for c in ws[1]:
            c.font = header_font
            c.fill = header_fill
            c.alignment = Alignment(horizontal="center")

        month_rows = [r for r in rows if r["jalali_year"] == y and r["jalali_month"] == m]
        for i, r in enumerate(month_rows, start=2):
            d = row_to_dict(r)
            values = [f"{r['jalali_month']:02d}/{r['jalali_day']:02d}", r["title"], r["tag_name"] or "",
                      r["start_time"] or "", r["end_time"] or "",
                      r["duration_minutes"] if r["duration_minutes"] is not None else "",
                      d["company_type"] or "", d["company_project"] or "", d["company_story"] or "",
                      r["fill_minutes"] or 0, "✓" if r["planned_flag"] else "", "✓" if r["outside_company"] else ""]
            for col, value in enumerate(values, start=1):
                ws.cell(row=i, column=col, value=value)
            for c in ws[i]:
                c.font = cell_font
                c.alignment = Alignment(horizontal="center")
        for idx, w in enumerate([16, 55, 20, 12, 12, 14, 14, 30, 45, 14, 10, 12], start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(idx)].width = w

    if not wb.sheetnames:
        wb.create_sheet("گزارش").append(["داده‌ای برای این بازه ثبت نشده است"])
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf
