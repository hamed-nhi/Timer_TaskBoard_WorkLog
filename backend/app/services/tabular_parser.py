"""Bounded Excel/CSV parsing into finished task records; no database or HTTP access."""
import csv
from datetime import date, datetime, time, timedelta
import io
import math
from pathlib import Path
import re
import zipfile

import openpyxl

from ..core.errors import ApiError
from ..core.time import g2j, j2g
from .timecalc import compute_duration
from .validation import clean_jalali, clean_time, clean_title

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
FIELDS = {
    'date': ('تاریخ', ['date','gregorian_date','jalali_date','تاریخ شمسی','تاریخ (ماه/روز)']),
    'title': ('شرح فعالیت', ['title','task','description','activity','عنوان','شرح','نام تسک']),
    'start': ('ساعت شروع', ['start','start_time','شروع','زمان شروع']),
    'end': ('ساعت پایان', ['end','end_time','پایان','زمان پایان']),
    'duration': ('مدت (دقیقه)', ['duration','duration_minutes','minutes','مدت','مدت زمان','زمان (دقیقه)']),
    'tag': ('تگ', ['tag','tag_name','برچسب']),
    'project': ('نام پروژه', ['project','project_name','پروژه']),
    'story': ('نام استوری', ['story','story_name','استوری']),
    'report_type': ('نوع گزارش', ['report_type','report type']),
    'fill': ('تکمیلی (دقیقه)', ['fill_minutes','fill','تکمیلی']),
    'planned': ('پلن شده', ['planned_flag','planned','برنامه‌ریزی‌شده']),
    'outside': ('خارج از شرکت', ['outside_company','outside']),
    'entered': ('ثبت در سامانه', ['entered_main_system','entered']),
}
DIGITS = str.maketrans('۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩يك', '01234567890123456789یک')


def text(value):
    return str(value if value is not None else '').strip()


def latin(value):
    return text(value).translate(DIGITS)


def header(value):
    return re.sub(r'[\s\u200c_\-()]+', '', latin(value)).casefold()


def number(value):
    if isinstance(value, bool):
        raise ApiError('عدد نامعتبر است')
    if isinstance(value, timedelta):
        result = value.total_seconds() / 60
    elif isinstance(value, time):
        result = value.hour * 60 + value.minute + value.second / 60
    else:
        raw = latin(value).replace('٫', '.')
        try:
            if re.fullmatch(r'\d+:\d{2}', raw):
                hours, minutes = map(int, raw.split(':'))
                if minutes > 59:
                    raise ValueError()
                result = hours * 60 + minutes
            else:
                result = float(raw)
        except (ValueError, TypeError, OverflowError):
            raise ApiError('مدت باید عدد دقیقه یا ساعت به شکل 02:30 باشد')
    if not math.isfinite(result) or not 0 <= result <= 10080:
        raise ApiError('مدت باید بین صفر و ۱۰۰۸۰ دقیقه باشد')
    return result


def clock(value):
    if isinstance(value, datetime):
        value = value.time()
    if isinstance(value, time):
        if value.second or value.microsecond:
            raise ApiError('ساعت دارای ثانیه است؛ آن را به HH:MM تبدیل کنید')
        return value.strftime('%H:%M')
    raw = latin(value)
    if not raw:
        return None
    if re.fullmatch(r'\d{1,2}:\d{2}', raw):
        raw = raw.zfill(5)
    return clean_time(raw)


def task_date(value, year=None):
    if isinstance(value, datetime):
        value = value.date()
    if isinstance(value, date):
        return value
    raw = latin(value)
    parts = re.fullmatch(r'(\d{1,4})[/-](\d{1,2})(?:[/-](\d{1,2}))?', raw)
    if not parts:
        raise ApiError('تاریخ باید YYYY/MM/DD شمسی یا YYYY-MM-DD میلادی باشد')
    a, b, c = parts.groups()
    try:
        if c is None:
            if not year:
                raise ApiError('تاریخ بدون سال است؛ سال فایل را مشخص کنید')
            return j2g(*clean_jalali(year, int(a), int(b)))
        if len(a) != 4:
            raise ValueError()
        if int(a) >= 1700:
            return date(int(a), int(b), int(c))
        return j2g(*clean_jalali(int(a), int(b), int(c)))
    except (ValueError, TypeError, OverflowError):
        raise ApiError('تاریخ نامعتبر است')


def flag(value):
    raw = latin(value).casefold()
    if raw in ('', '0', 'false', 'no', 'خیر', 'نه'):
        return False
    if raw in ('1', 'true', 'yes', 'بله', '✓', '✔', 'x'):
        return True
    raise ApiError('مقدار بله/خیر نامعتبر است')


def read_file(content, filename):
    if not content or len(content) > MAX_BYTES:
        raise ApiError('فایل باید خالی نباشد و حداکثر ۵ مگابایت باشد')
    suffix = Path(filename or '').suffix.lower()
    if suffix == '.csv':
        decoded = None
        encodings = ('utf-16',) if content.startswith((b'\xff\xfe', b'\xfe\xff')) else ('utf-8-sig', 'cp1256')
        for encoding in encodings:
            try:
                decoded = content.decode(encoding)
                break
            except UnicodeError:
                continue
        if decoded is None or '\x00' in decoded:
            raise ApiError('کدگذاری CSV نامعتبر است؛ فایل را با UTF-8 ذخیره کنید')
        try:
            dialect = csv.Sniffer().sniff(decoded[:8192], delimiters=',;\t')
        except csv.Error:
            dialect = csv.excel
        csv.field_size_limit(MAX_BYTES)
        try:
            reader = csv.reader(io.StringIO(decoded, newline=''), dialect)
            rows = []
            for row in reader:
                if len(rows) > MAX_ROWS:
                    raise ApiError('حداکثر ۵۰۰۰ سطر در هر فایل مجاز است')
                if len(row) > 100:
                    raise ApiError('حداکثر ۱۰۰ ستون مجاز است')
                rows.append(row)
            return [('CSV', rows)]
        except csv.Error:
            raise ApiError('ساختار CSV نامعتبر است')
    if suffix != '.xlsx':
        raise ApiError('فقط فایل XLSX یا CSV پذیرفته می‌شود؛ فایل XLS را ابتدا به XLSX تبدیل کنید')
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ApiError('حجم بازشدهٔ فایل اکسل بیش از حد مجاز است')
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if len(workbook.worksheets) > 24:
                raise ApiError('حداکثر ۲۴ برگه در هر فایل مجاز است')
            sheets, total = [], 0
            for sheet in workbook.worksheets:
                if sheet.max_column and sheet.max_column > 100:
                    raise ApiError('حداکثر ۱۰۰ ستون مجاز است')
                rows = []
                for cells in sheet.iter_rows():
                    total += 1
                    if total > MAX_ROWS + 24:
                        raise ApiError('حداکثر ۵۰۰۰ سطر در هر فایل مجاز است')
                    values = []
                    for cell in cells:
                        if cell.data_type == 'f':
                            raise ApiError('فایل فرمول دارد؛ ابتدا مقدار سلول‌ها را به‌صورت Values ذخیره کنید')
                        values.append(cell.value)
                    rows.append(values)
                sheets.append((sheet.title, rows))
            return sheets
        finally:
            workbook.close()
    except ApiError:
        raise
    except Exception:
        raise ApiError('فایل اکسل خراب، رمزدار یا نامعتبر است')


def parse_file(content, filename, mapping=None, default_year=None):
    sheets = read_file(content, filename)
    if not isinstance(mapping, dict) or any(k not in FIELDS or not isinstance(v, str) for k, v in mapping.items()):
        if mapping is not None:
            raise ApiError('نگاشت ستون‌ها نامعتبر است')
        mapping = {}
    if default_year not in (None, ''):
        clean_jalali(default_year, 1, 1)
    rows, errors, warnings, columns, chosen = [], [], [], [], {}
    count = 0
    for sheet_name, raw_rows in sheets:
        raw_rows = list(raw_rows)
        leading_rows = 0
        while raw_rows and not any(text(v) for v in raw_rows[0]):
            raw_rows.pop(0)
            leading_rows += 1
        if not raw_rows or not any(text(v) for row in raw_rows for v in row):
            continue
        names = [text(v) for v in raw_rows[0]]
        if not columns:
            columns = names
            nonempty = [header(n) for n in names if n]
            if len(set(nonempty)) != len(nonempty):
                raise ApiError('نام ستون‌ها نباید تکراری باشد')
            for key, (label, aliases) in FIELDS.items():
                if key in mapping:
                    chosen[key] = mapping[key]
                else:
                    chosen[key] = next((n for n in names if header(n) in {header(label), *(header(a) for a in aliases)}), '')
            if any(v and v not in names for v in chosen.values()):
                raise ApiError('ستون انتخاب‌شده در فایل وجود ندارد')
        elif names != columns:
            raise ApiError('ستون‌های برگه‌های اکسل یکسان نیستند؛ هر قالب را جدا وارد کنید')
        year = default_year
        match = re.search(r'(?<!\d)(1[0-6]\d{2})(?!\d)', latin(sheet_name))
        if not year and match:
            year = int(match[1])
        if not chosen.get('date') or not chosen.get('title'):
            errors.append({'sheet': sheet_name, 'row': 1 + leading_rows, 'message': 'ستون تاریخ و شرح فعالیت را انتخاب کنید'})
            continue
        for index, raw in enumerate(raw_rows[1:], start=2 + leading_rows):
            if not any(text(v) for v in raw):
                continue
            count += 1
            if count > MAX_ROWS:
                raise ApiError('حداکثر ۵۰۰۰ سطر در هر فایل مجاز است')
            values = dict(zip(names, raw))
            value = lambda field: values.get(chosen.get(field, ''))  # noqa: E731
            try:
                day = task_date(value('date'), year)
                start, end = clock(value('start')), clock(value('end'))
                if bool(start) != bool(end):
                    raise ApiError('ساعت شروع و پایان را هر دو وارد کنید یا هر دو را خالی بگذارید')
                duration = number(value('duration')) if text(value('duration')) else None
                if duration is None:
                    if not start or not end:
                        raise ApiError('مدت یا هر دو ساعت شروع و پایان الزامی است')
                    duration = compute_duration(start, end)
                row = {'date': day.isoformat(), 'title': clean_title(value('title')), 'start': start, 'end': end,
                       'duration': duration, 'fill': number(value('fill')) if text(value('fill')) else 0,
                       'planned': flag(value('planned')), 'outside': flag(value('outside')), 'entered': flag(value('entered'))}
                for field in ('tag', 'project', 'story', 'report_type'):
                    row[field] = text(value(field))
                    if len(row[field]) > (60 if field == 'tag' else 150):
                        raise ApiError('نام تگ حداکثر ۶۰ و نام پروژه، استوری و نوع گزارش حداکثر ۱۵۰ کاراکتر است')
                if row['story'] and not row['project']:
                    raise ApiError('نام پروژه برای استوری الزامی است')
                if row['project'] and not row['story']:
                    row['story'] = 'بدون استوری'
                    warnings.append({'sheet': sheet_name, 'row': index, 'message': 'پروژه فاقد استوری است؛ در استوری «بدون استوری» قرار می‌گیرد'})
                if (row['project'] or row['report_type']) and not row['tag']:
                    row['tag'] = row['story'] or row['report_type']
                    if len(row['tag']) > 60:
                        raise ApiError('برای این ردیف یک نام تگ حداکثر ۶۰ کاراکتری مشخص کنید')
                if start and end and abs(duration - compute_duration(start, end)) > .01:
                    warnings.append({'sheet': sheet_name, 'row': index, 'message': 'مدت ثبت‌شده با فاصلهٔ شروع و پایان فرق دارد؛ همان مدت فایل حفظ می‌شود'})
                rows.append({**row, 'sheet': sheet_name, 'row': index})
            except ApiError as error:
                errors.append({'sheet': sheet_name, 'row': index, 'message': str(error)})
    if not columns:
        raise ApiError('فایل داده‌ای ندارد')
    return {'columns': columns, 'mapping': chosen, 'fields': [{'key': k, 'label': v[0]} for k,v in FIELDS.items()],
            'rows': rows, 'errors': errors, 'warnings': warnings, 'total_rows': count}
