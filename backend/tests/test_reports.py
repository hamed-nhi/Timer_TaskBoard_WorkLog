"""Dashboard comparison, weekly timesheet, calendar, goals, entry report and the Excel export."""
import io
from datetime import date, timedelta

import jdatetime
import openpyxl

from conftest import sql
from test_tags_planned import done
from app.routers import reports


def freeze_today(monkeypatch, y, m, d):
    monkeypatch.setattr(reports, "today_jalali", lambda: jdatetime.date(y, m, d))


def dash(client, **params):
    qs = "&".join(f"{k}={v}" for k, v in params.items())
    return client.get(f"/api/dashboard?{qs}").get_json()


# ---- dashboard comparison over the same number of days ----

def test_running_month_is_compared_with_the_same_first_days_of_the_previous_month(client, monkeypatch):
    freeze_today(monkeypatch, 1405, 7, 10)
    done(client, 120, 1405, 7, 3)
    done(client, 60, 1405, 7, 10)
    done(client, 100, 1405, 6, 2)
    done(client, 50, 1405, 6, 10)
    done(client, 900, 1405, 6, 11)
    done(client, 900, 1405, 6, 25)
    c = dash(client, year=1405, month=7)["comparison"]
    assert (c["days"], c["current_minutes"], c["prev_minutes"], c["diff_minutes"], c["diff_percent"]) == (10, 180, 150, 30, 20.0)
    assert c["current_label"] == "1 تا 10 مهر" and c["prev_label"] == "1 تا 10 شهریور"


def test_a_finished_month_is_compared_in_full(client, monkeypatch):
    freeze_today(monkeypatch, 1405, 9, 5)
    done(client, 60, 1405, 7, 29)
    done(client, 30, 1405, 6, 30)
    done(client, 500, 1405, 6, 31)
    c = dash(client, year=1405, month=7)["comparison"]
    assert c["days"] == 30 and c["current_minutes"] == 60 and c["prev_minutes"] == 30


def test_first_month_of_the_year_and_first_day_label(client, monkeypatch):
    freeze_today(monkeypatch, 1405, 1, 4)
    done(client, 40, 1405, 1, 2)
    done(client, 25, 1404, 12, 3)
    done(client, 500, 1404, 12, 20)
    c = dash(client, year=1405, month=1)["comparison"]
    assert (c["days"], c["current_minutes"], c["prev_minutes"]) == (4, 40, 25)
    freeze_today(monkeypatch, 1405, 7, 1)
    c = dash(client, year=1405, month=7)["comparison"]
    assert c["days"] == 1 and c["current_label"] == "1 مهر"


def test_no_previous_data_gives_no_percent(client, monkeypatch):
    freeze_today(monkeypatch, 1405, 7, 10)
    done(client, 60, 1405, 7, 2)
    c = dash(client, year=1405, month=7)["comparison"]
    assert c["prev_minutes"] == 0 and c["diff_percent"] is None and c["diff_minutes"] == 60


def test_running_and_finished_year(client, monkeypatch):
    freeze_today(monkeypatch, 1405, 7, 10)
    done(client, 100, 1405, 2, 5)
    done(client, 10, 1405, 7, 10)
    done(client, 80, 1404, 2, 5)
    done(client, 5, 1404, 7, 10)
    done(client, 600, 1404, 7, 11)
    done(client, 600, 1404, 12, 1)
    c = dash(client, year=1405)["comparison"]
    assert (c["current_minutes"], c["prev_minutes"]) == (110, 85)
    assert "10 مهر 1405" in c["current_label"] and "10 مهر 1404" in c["prev_label"]
    done(client, 60, 1404, 12, 28)
    done(client, 30, 1403, 12, 28)
    c = dash(client, year=1404)["comparison"]
    assert (c["current_minutes"], c["prev_minutes"]) == (60 + 600 + 80 + 5 + 600, 30)


def test_future_period_has_no_comparison_and_open_tasks_are_ignored(client, monkeypatch):
    freeze_today(monkeypatch, 1405, 7, 10)
    assert dash(client, year=1405, month=8)["comparison"] is None
    assert dash(client, year=1406)["comparison"] is None
    assert dash(client)["comparison"] is None
    client.post("/api/tasks", json={"title": "در جریان", "jalali_year": 1405, "jalali_month": 7,
                                    "jalali_day": 2, "start_time": "09:00"})
    assert dash(client, year=1405, month=7)["comparison"]["current_minutes"] == 0


def test_dashboard_payload(client):
    a = client.post("/api/tags", json={"name": "الف"}).get_json()
    done(client, 60, 1405, 7, 1, a["id"])
    done(client, 30, 1405, 7, 1, None)
    done(client, 200, 1405, 7, 2, a["id"])
    d = dash(client, year=1405, month=7)
    assert d["active_days"] == 2 and d["best_day"] == {"date": "1405/07/02", "minutes": 200}
    assert d["daily_goal_minutes"] > 0 and d["days_in_month"] == 30
    by_tag = {t["tag_name"]: t for t in d["by_tag"]}
    assert by_tag["الف"]["tag_id"] == a["id"] and by_tag["الف"]["minutes"] == 260 and by_tag["الف"]["count"] == 2
    assert by_tag["بدون تگ"]["tag_id"] is None and by_tag["بدون تگ"]["minutes"] == 30
    assert d["by_tag"][0]["tag_name"] == "الف" and d["by_month"][0]["month_name"] == "مهر"


def test_empty_dashboard_does_not_crash(client):
    d = dash(client, year=1399, month=1)
    assert d["best_day"] is None and d["active_days"] == 0 and d["total_minutes"] == 0


# ---- weekly timesheet ----

def j(d):
    x = jdatetime.date.fromgregorian(date=d)
    return x.year, x.month, x.day


def put(client, minutes, d, tag_id=None):
    y, m, dd = j(d)
    return done(client, minutes, y, m, dd, tag_id=tag_id)


def ts(client, d=None):
    return client.get("/api/timesheet" + (f"?date={d.isoformat()}" if d else "")).get_json()


SAT = date(2026, 10, 3)


def test_week_runs_saturday_to_friday_whatever_day_is_asked(client):
    for offset in range(7):
        data = ts(client, SAT + timedelta(days=offset))
        assert data["week_start"] == SAT.isoformat() and data["week_end"] == (SAT + timedelta(days=6)).isoformat()
        assert data["days"][0]["weekday"] == "شنبه" and data["days"][-1]["weekday"] == "جمعه"


def test_cells_rows_and_totals(client):
    a = client.post("/api/tags", json={"name": "الف"}).get_json()["id"]
    b = client.post("/api/tags", json={"name": "ب"}).get_json()["id"]
    put(client, 90, SAT, a)
    put(client, 30, SAT, a)
    put(client, 60, SAT + timedelta(days=2), a)
    put(client, 45, SAT + timedelta(days=2), b)
    put(client, 20, SAT + timedelta(days=6))
    data = ts(client, SAT)
    rows = {r["tag_name"]: r for r in data["rows"]}
    assert rows["الف"]["minutes"] == [120, 0, 60, 0, 0, 0, 0] and rows["الف"]["total"] == 180
    assert rows["ب"]["minutes"][2] == 45 and rows["بدون تگ"]["minutes"][6] == 20
    assert data["day_totals"] == [120, 0, 105, 0, 0, 0, 20] and data["total"] == 245
    assert [r["tag_name"] for r in data["rows"]] == ["الف", "ب", "بدون تگ"]


def test_other_weeks_excluded_and_navigation_links(client):
    put(client, 50, SAT - timedelta(days=1))
    put(client, 70, SAT + timedelta(days=7))
    data = ts(client, SAT)
    assert data["rows"] == [] and data["total"] == 0
    assert data["prev"] == (SAT - timedelta(days=7)).isoformat() and data["next"] == (SAT + timedelta(days=7)).isoformat()
    assert ts(client, date.fromisoformat(data["next"]))["total"] == 70


def test_timesheet_labels_bad_date_and_read_only(client):
    data = ts(client)
    assert len(data["days"]) == 7 and sum(1 for d in data["days"] if d["is_today"]) == 1
    first = data["days"][0]
    assert (first["jalali_year"], first["jalali_month"], first["jalali_day"]) == j(date.fromisoformat(first["gregorian"]))
    assert client.get("/api/timesheet?date=not-a-date").status_code == 400
    put(client, 33, SAT)
    before = sql("SELECT * FROM tasks")
    ts(client, SAT)
    assert sql("SELECT * FROM tasks") == before


# ---- calendar / goals / entry report ----

def test_calendar_month(client):
    done(client, 60, 1405, 7, 3)
    client.post("/api/planned", json={"title": "برنامه", "jalali_year": 1405, "jalali_month": 7, "jalali_day": 3})
    cal = client.get("/api/calendar?year=1405&month=7").get_json()
    assert cal["days_in_month"] == 30 and len(cal["days"]) == 30 and cal["month_name"] == "مهر"
    d3 = cal["days"][2]
    assert d3["minutes"] == 60 and d3["task_count"] == 1 and d3["planned_count"] == 1
    assert client.get("/api/calendar?year=1405&month=13").status_code == 400
    nowruz = client.get("/api/calendar?year=1405&month=1").get_json()["days"][0]
    assert nowruz["holiday"] is True


def test_goals_and_progress(client):
    assert client.post("/api/settings", json={"daily_goal_minutes": 0}).status_code == 400
    assert client.post("/api/settings", json={"daily_goal_minutes": "x"}).status_code == 400
    assert client.post("/api/settings", json={"google_sheet_url": "https://evil.example/x"}).status_code == 400
    s = client.post("/api/settings", json={"daily_goal_minutes": 300, "weekly_goal_minutes": 1500,
                                           "google_sheet_url": "https://docs.google.com/spreadsheets/d/abc"}).get_json()
    assert s["daily_goal_minutes"] == 300 and s["google_sheet_url"].endswith("abc")
    j_ = client.get("/api/today").get_json()
    done(client, 60, j_["jalali_year"], j_["jalali_month"], j_["jalali_day"])
    p = client.get("/api/goal-progress").get_json()
    assert p["today_minutes"] == 60 and p["week_minutes"] >= 60 and p["weekly_goal_minutes"] == 1500


def test_entry_report(client):
    t = done(client, 30, 1405, 7, 1)
    done(client, 30, 1405, 7, 1)
    client.post(f"/api/tasks/{t['id']}/entered", json={"entered": True})
    r = client.get("/api/entry-report?year=1405&month=7").get_json()
    assert r == [{"date": "1405/07/01", "total_tasks": 2, "entered_tasks": 1, "total_minutes": 60.0, "all_entered": False}]


# ---- Excel export ----

def sheet_titles(res):
    wb = openpyxl.load_workbook(io.BytesIO(res.data))
    return wb, [row[1] for ws in wb for row in ws.iter_rows(min_row=2, values_only=True) if row[1]]


def test_export_empty_still_returns_a_valid_workbook(client):
    res = client.get("/api/export?year=1405")
    assert res.status_code == 200 and "spreadsheetml" in res.headers["content-type"]
    assert len(openpyxl.load_workbook(io.BytesIO(res.data)).sheetnames) >= 1


def test_export_sheets_headers_and_year_filter(client):
    done(client, 60, 1405, 7, 3, title="جلسه اول")
    done(client, 30, 1405, 8, 1, title="جلسه دوم")
    done(client, 30, 1404, 7, 3, title="سال قبل")
    res = client.get("/api/export?year=1405")
    wb, titles = sheet_titles(res)
    assert "مهر 1405" in wb.sheetnames and "آبان 1405" in wb.sheetnames
    assert "سال قبل" not in titles and "جلسه اول" in titles
    ws = wb["مهر 1405"]
    header = [c.value for c in ws[1]]
    assert header[:6] == ["تاریخ (ماه/روز)", "شرح فعالیت", "تگ", "ساعت شروع", "ساعت پایان", "مدت (دقیقه)"]
    assert header[6:] == ["نوع گزارش", "نام پروژه", "نام استوری", "تکمیلی (دقیقه)", "پلن شده", "خارج از شرکت"]
    row = [c.value for c in ws[2]]
    assert row[0] == "07/03" and row[1] == "جلسه اول" and row[3] == "08:00" and row[5] == 60.0
    assert "سال قبل" in sheet_titles(client.get("/api/export"))[1]


def test_export_running_task_has_blank_end_and_duration(client):
    client.post("/api/tasks", json={"title": "در حال انجام", "jalali_year": 1405, "jalali_month": 7,
                                    "jalali_day": 3, "start_time": "09:00"})
    row = [c.value for c in openpyxl.load_workbook(io.BytesIO(client.get("/api/export?year=1405").data))["مهر 1405"][2]]
    assert row[1] == "در حال انجام" and row[4] in (None, "") and row[5] in (None, "")


def test_export_has_company_columns(client, admin):
    t_id = admin.post("/api/admin/report-types", json={"name": "گزارش فنی"}).get_json()["id"]
    pid = admin.post("/api/admin/projects", json={"name": "پروژه"}).get_json()["id"]
    sid = admin.post(f"/api/admin/projects/{pid}/stories", json={"name": "استوری"}).get_json()["id"]
    tag = client.post("/api/tags", json={"name": "تگ"}).get_json()
    client.put(f"/api/tags/{tag['id']}/company", json={"report_type_id": t_id, "story_id": sid})
    t = done(client, 60, 1405, 7, 3, tag_id=tag["id"])
    client.post(f"/api/tasks/{t['id']}/company", json={"planned_flag": True})
    row = [c.value for c in openpyxl.load_workbook(io.BytesIO(client.get("/api/export?year=1405").data))["مهر 1405"][2]]
    assert row[6:9] == ["گزارش فنی", "پروژه", "استوری"] and row[10] == "✓" and row[11] in (None, "")


def test_my_backup_download(client):
    done(client, 10, 1405, 7, 1)
    res = client.get("/api/me/backup")
    assert "attachment" in res.headers["content-disposition"]
    data = res.json()
    assert len(data["tasks"]) == 1 and len(data["task_segments"]) == 1 and data["user"]["username"] == "alice"
